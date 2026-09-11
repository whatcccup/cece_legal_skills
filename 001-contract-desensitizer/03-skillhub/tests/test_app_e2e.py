#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""端到端回归：Web 应用走完「上传 → 复核页 → 脱敏 → 下载 → 还原」全链路。

重点验证三件事：
  1. 复核页展示的 [标签#N] 与最终 docx 里写的**完全一致**；
  2. 同一个人/企业（含简称）在整篇里只有一个编号；
  3. 还原后与原文逐段一致。

跑法：
    /Users/agent/.workbuddy/binaries/python/envs/default/bin/python tests/test_app_e2e.py
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import urllib.request
import uuid
import zipfile
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import contract_app_server as S  # noqa: E402
from docx import Document  # noqa: E402

FAILS: list = []
HOST, PORT = "127.0.0.1", 18899
BASE = f"http://{HOST}:{PORT}"

# 同时覆盖两类漏检：带内部分组空格的座机、校验码错误的统一社会信用代码
SAMPLE = """甲方：北京星海智能科技有限公司（以下简称“星海智能”）
统一社会信用代码：91110108MA8A2X5K3R
联系电话：010-8666 2188
法定代表人：张三
乙方：上海星光数据服务有限公司
授权代表：李四
第一条 星海智能 委托乙方提供数据服务，乙方指派张三 作为项目负责人。
第二条 张三 应于每季度向星海智能 提交报告，同时抄送李四。
第三条 本协议由北京星海智能科技有限公司 与上海星光数据服务有限公司 共同签署。
第四条 星海智能 的联系电话为 010-86662188。
第五条 甲方法定代表人张三 签字，乙方授权代表李四 签字。
"""


def check(name: str, ok: bool, extra: str = "") -> None:
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  {extra}" if extra and not ok else ""))
    if not ok:
        FAILS.append(name)


def http_get(path: str) -> bytes:
    with urllib.request.urlopen(BASE + path, timeout=120) as r:
        return r.read()


def http_post_json(path: str, payload: dict) -> dict:
    req = urllib.request.Request(
        BASE + path, data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=600) as r:
        return json.loads(r.read().decode("utf-8"))


def http_upload(docx_path: str, mode: str = "redact") -> dict:
    boundary = "----wb" + uuid.uuid4().hex
    with open(docx_path, "rb") as f:
        data = f.read()
    body = io.BytesIO()
    for name, val in (("mode", mode),):
        body.write(f"--{boundary}\r\n".encode())
        body.write(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        body.write(val.encode() + b"\r\n")
    body.write(f"--{boundary}\r\n".encode())
    body.write(b'Content-Disposition: form-data; name="file"; filename="input.docx"\r\n')
    body.write(b"Content-Type: application/vnd.openxmlformats-officedocument"
               b".wordprocessingml.document\r\n\r\n")
    body.write(data + b"\r\n")
    body.write(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        BASE + "/api/upload", data=body.getvalue(),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST")
    with urllib.request.urlopen(req, timeout=600) as r:
        return json.loads(r.read().decode("utf-8"))


def docx_text(path: str) -> str:
    """取 docx 全部文本（含表格、页眉页脚）。"""
    with zipfile.ZipFile(path) as z:
        xml = "".join(z.read(n).decode("utf-8") for n in z.namelist()
                      if n.startswith("word/") and n.endswith(".xml"))
    return "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S))


def placeholder_nums(text: str, label: str) -> list:
    return sorted({int(x) for x in re.findall(rf"\[{re.escape(label)}#(\d+)\]", text)})


def run() -> int:
    tmp = tempfile.mkdtemp(prefix="app_e2e_")
    src = os.path.join(tmp, "input.docx")
    d = Document()
    for line in SAMPLE.splitlines():
        d.add_paragraph(line)
    d.save(src)

    server = S.AppHandler and S.start_app_server(HOST, PORT, os.path.join(tmp, "sessions"))
    th = threading.Thread(target=server.serve_forever, daemon=True)
    th.start()
    time.sleep(0.5)
    try:
        print("== 1. 上传 + 复核页编号 ==")
        up = http_upload(src)
        check("上传成功", up.get("ok") is True, str(up))
        sid = up["sid"]

        ws = http_get(f"/workspace?sid={sid}").decode("utf-8")
        check("工作台可打开", len(ws) > 500)
        rv = http_get(f"/review?sid={sid}").decode("utf-8")
        check("左右对比页可打开", len(rv) > 1000 and "var DATA" in rv)

        # 复核页里的编号（HTML 内嵌 JSON）
        ws_nums = defaultdict(list)
        for lab, n in re.findall(r"\[(自然人姓名|组织机构名称)#(\d+)\]", rv):
            ws_nums[lab].append(int(n))
        ws_sets = {k: sorted(set(v)) for k, v in ws_nums.items()}
        print(f"     复核页编号：{ws_sets}")

        print("== 2. 执行脱敏 ==")
        ap = http_post_json(f"/api/apply/{sid}", {"enabled_types": []})
        check("脱敏成功", ap.get("ok") is True, str(ap))
        out_name = ap["output_name"]
        out_path = os.path.join(tmp, "sessions", sid, "out", out_name)
        check("生成脱敏文件", os.path.exists(out_path), out_path)

        import urllib.parse
        dl = http_get(f"/download/{sid}/{urllib.parse.quote(out_name)}")
        check("下载成功", len(dl) > 1000)
        with open(out_path, "wb") as f:
            f.write(dl)

        text = docx_text(out_path)
        doc_json = defaultdict(list)
        for lab, n in re.findall(r"\[(自然人姓名|组织机构名称)#(\d+)\]", text):
            doc_json[lab].append(int(n))
        doc_sets = {k: sorted(set(v)) for k, v in doc_json.items()}
        print(f"     产物编号  ：{doc_sets}")

        print("== 3. 复核页编号 == 产物编号 ==")
        check("自然人姓名编号一致", ws_sets.get("自然人姓名") == doc_sets.get("自然人姓名"),
              f"{ws_sets.get('自然人姓名')} vs {doc_sets.get('自然人姓名')}")
        check("组织机构名称编号一致", ws_sets.get("组织机构名称") == doc_sets.get("组织机构名称"),
              f"{ws_sets.get('组织机构名称')} vs {doc_sets.get('组织机构名称')}")

        print("== 4. 编号连续性 & 同名同号 ==")
        check("人名编号为 1..N 连续", doc_sets.get("自然人姓名") == [1, 2],
              str(doc_sets.get("自然人姓名")))
        check("机构编号为 1..N 连续", doc_sets.get("组织机构名称") == [1, 2],
              str(doc_sets.get("组织机构名称")))
        check("张三全部为同一编号", text.count("[自然人姓名#1]") == SAMPLE.count("张三"),
              f"count={text.count('[自然人姓名#1]')} vs {SAMPLE.count('张三')}")
        check("全称与简称共用同一编号（共 6 处）",
              text.count("[组织机构名称#1]") == 6, f"count={text.count('[组织机构名称#1]')}")

        print("== 5. 漏检修复 ==")
        check("座机 010-8666 2188 被识别", "010-8666 2188" not in text)
        check("错误校验码的统一社会信用代码被识别", "91110108MA8A2X5K3R" not in text)

        print("== 6. 还原 ==")
        rs = http_post_json(f"/api/restore/{sid}", {})
        check("还原成功", rs.get("ok") is True, str(rs))
        back_path = os.path.join(tmp, "sessions", sid, "out", rs.get("output_name", ""))
        check("生成还原文件", os.path.exists(back_path), back_path)
        a = [p.text for p in Document(src).paragraphs]
        b = [p.text for p in Document(back_path).paragraphs]
        check("还原后逐段一致", a == b)
        if a != b:
            for i, (x, y) in enumerate(zip(a, b)):
                if x != y:
                    print(f"      第{i}段\n        原={x!r}\n        还={y!r}")
                    break
    finally:
        try:
            server.shutdown()
            server.server_close()
        except Exception:
            pass
        shutil.rmtree(tmp, ignore_errors=True)

    print()
    if FAILS:
        print(f"✘ 失败 {len(FAILS)} 项：" + "；".join(FAILS))
        return 1
    print("✔ 全链路通过")
    return 0


if __name__ == "__main__":
    sys.exit(run())
