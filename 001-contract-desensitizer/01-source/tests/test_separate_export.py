"""Word contains placeholders only; sensitive originals live exclusively in JSON."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
import contract_redactor as R


class SeparateExport(unittest.TestCase):
    def test_external_docx_and_separate_mapping_restore(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / '虚构合同.docx'
            doc = Document(); doc.add_paragraph('联系电话：13800138000'); doc.save(src)
            out = Path(tmp) / '虚构合同_脱敏.docx'
            result = R.redact_file(str(src), types=['PHONE_MOBILE'], out_path=str(out), export_mode='external')
            self.assertIsNotNone(result['mapping'], '必须同时生成单独的 JSON')
            self.assertFalse(result['mapping_embedded'])
            for options in ({'restore_mode': True}, {'export_mode': 'reversible'}):
                legacy = R.redact_file(str(src), types=['PHONE_MOBILE'], out_path=str(Path(tmp)/'legacy.docx'), **options)
                self.assertFalse(legacy['mapping_embedded'])
                with zipfile.ZipFile(legacy['output']) as z:
                    self.assertFalse(any(b'13800138000' in z.read(n) for n in z.namelist()))
            self.assertIsNone(result['log'])
            with zipfile.ZipFile(out) as z:
                self.assertFalse(any(n.startswith('customXml/') for n in z.namelist()))
                self.assertFalse(any(b'13800138000' in z.read(n) for n in z.namelist()))
            mapping = json.loads(Path(result['mapping']).read_text())
            self.assertEqual(mapping['items'][0]['original'], '13800138000')
            restored = Path(tmp) / '还原.docx'
            R.restore_file(str(out), mapping_path=result['mapping'], out_path=str(restored))
            self.assertEqual(Document(restored).paragraphs[0].text, doc.paragraphs[0].text)


if __name__ == '__main__': unittest.main()
