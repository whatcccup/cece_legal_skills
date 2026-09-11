#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""回归测试：同一实体编号一致 + 座机/统一社会信用代码漏检修复。

跑法：
    /Users/agent/.workbuddy/binaries/python/envs/default/bin/python tests/test_entity_consistency.py
"""
from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import contract_sensitive_detector as D  # noqa: E402
import contract_redactor as R  # noqa: E402
from docx import Document  # noqa: E402

FAILS: list = []


def check(name: str, ok: bool, extra: str = "") -> None:
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  {extra}" if extra and not ok else ""))
    if not ok:
        FAILS.append(name)


SAMPLE = """甲方：北京星海智能科技有限公司（以下简称“星海智能”）
统一社会信用代码：91110108MA8A2X5K3R
住所：北京市海淀区中关村南大街 5 号院 3 号楼 1201
联系电话：010-8666 2188
法定代表人：张三

乙方：上海星光数据服务有限公司
统一社会信用代码：91310115MA7B6Q2N8U
联系电话：010-8666 2188
授权代表：李四

第一条 星海智能 委托乙方提供数据服务，乙方指派张三 作为项目负责人。
第二条 张三 应于每季度向星海智能 提交报告，同时抄送李四。
第三条 本协议由北京星海智能科技有限公司 与上海星光数据服务有限公司 共同签署。
第四条 星海智能 的联系电话为 010-86662188，乙方联系电话为 010-8666 2188。
第五条 甲方法定代表人张三 签字，乙方授权代表李四 签字。
"""


def build_docx(path: str) -> None:
    doc = Document()
    for line in SAMPLE.splitlines():
        doc.add_paragraph(line)
    doc.save(path)


def run() -> int:
    tmp = tempfile.mkdtemp(prefix="entity_consistency_")
    try:
        src = os.path.join(tmp, "sample.docx")
        build_docx(src)

        settings = D.Settings(min_confidence=0.5, min_severity="low")
        doc = D.load_document(src)
        engine = D.DetectorEngine(settings)
        cands = engine.scan_document(doc)
        report = D.build_report(doc, cands, settings)

        print("== 1. 座机识别（含内部分组空格） ==")
        landlines = [it["value"] for it in report["items"] if it["type"] == "PHONE_LANDLINE"]
        check("识别出 010-8666 2188", any("8666" in v and "2188" in v for v in landlines),
              f"实际={landlines}")

        print("== 2. 统一社会信用代码识别（校验码错误也要识别） ==")
        usccs = [it["value"] for it in report["items"] if it["type"] == "USCC"]
        check("识别 91110108MA8A2X5K3R（校验错）", "91110108MA8A2X5K3R" in usccs, f"实际={usccs}")
        check("识别 91310115MA7B6Q2N8U", "91310115MA7B6Q2N8U" in usccs, f"实际={usccs}")

        print("== 3. 实体编号：同一实体必须同号 ==")
        numberer = D.EntityNumberer(report["items"])
        p_zhang = numberer.number("PERSON_NAME", "张三")
        p_li = numberer.number("PERSON_NAME", "李四")
        o_full = numberer.number("ORG_NAME", "北京星海智能科技有限公司")
        o_short = numberer.number("ORG_NAME", "星海智能")
        o_other = numberer.number("ORG_NAME", "上海星光数据服务有限公司")
        check("张三 / 李四 编号不同", p_zhang != p_li, f"{p_zhang} vs {p_li}")
        check("全称与简称「星海智能」同号", o_full == o_short, f"{o_full} vs {o_short}")
        check("甲方与乙方 编号不同", o_full != o_other, f"{o_full} vs {o_other}")

        print("== 4. 脱敏产物：同一个人的所有占位符完全相同 ==")
        out = os.path.join(tmp, "sample_脱敏.docx")
        all_types = sorted({it["type"] for it in report["items"]})
        R.redact_file(src, types=all_types, out_path=out, restore_mode=True)
        redacted = Document(out)
        text = "\n".join(p.text for p in redacted.paragraphs)

        person_phs = set(re.findall(r"\[自然人姓名#(\d+)\]", text))
        org_phs = set(re.findall(r"\[组织机构名称#(\d+)\]", text))
        check("人名占位符只有 2 种（张三 / 李四）", person_phs == {"1", "2"}, f"实际={sorted(person_phs)}")
        check("机构占位符只有 2 种（甲乙双方）", org_phs == {"1", "2"}, f"实际={sorted(org_phs)}")

        zhang_ph = f"[自然人姓名#{p_zhang}]"
        check("张三的占位符处处相同", text.count(zhang_ph) == SAMPLE.count("张三"),
              f"次数={text.count(zhang_ph)} vs 原文 {SAMPLE.count('张三')}")

        # 全称出现 2 次 + 简称出现 4 次，必须都是 [组织机构名称#1]
        org1_ph = f"[组织机构名称#{o_full}]"
        check("全称与简称共用同一个占位符（共 6 处）", text.count(org1_ph) == 6,
              f"实际={text.count(org1_ph)}")
        check("「星海智能」不再单独占号", "[组织机构名称#" not in text or
              sorted(org_phs) == ["1", "2"], f"实际={sorted(org_phs)}")
        print("     脱敏片段：" + text.splitlines()[0])

        print("== 5. 还原：必须与原文完全一致 ==")
        back = os.path.join(tmp, "sample_还原.docx")
        rr = R.restore_file(out, out_path=back)
        a = "\n".join(p.text for p in Document(src).paragraphs)
        b = "\n".join(p.text for p in Document(back).paragraphs)
        check("还原后逐段一致", a == b)
        if a != b:
            for i, (x, y) in enumerate(zip(a.splitlines(), b.splitlines())):
                if x != y:
                    print(f"      第{i}段 原={x!r} 还原={y!r}")
                    break
        check("还原覆盖率 100%", rr.get("restore_ratio", 1) == 1 or not rr.get("missing"),
              f"missing={rr.get('missing')}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print()
    if FAILS:
        print(f"✘ 失败 {len(FAILS)} 项：" + "；".join(FAILS))
        return 1
    print("✔ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(run())
