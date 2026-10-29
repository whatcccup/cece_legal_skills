"""诉讼当事人信息漏检回归；全部使用虚构内容。"""
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import contract_sensitive_detector as D
import contract_redactor as R
from docx import Document

class LitigationDetection(unittest.TestCase):
    def scan(self, text):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'fictional.docx'
            doc = Document()
            doc.add_paragraph(text)
            doc.save(path)
            settings = D.Settings(min_confidence=0.5, min_severity='low')
            loaded = D.load_document(str(path))
            candidates = D.DetectorEngine(settings).scan_document(loaded)
            return D.build_report(loaded, candidates, settings)['items']

    def test_body_names_without_role_labels(self):
        text = '赵星辰向朱宁支付款项。欧阳晓明收到转账后通知钱晓。赵星辰随后离开。'
        names = [i['value'] for i in self.scan(text) if i['type'] == 'PERSON_NAME']
        for name in ('赵星辰', '朱宁', '欧阳晓明', '钱晓'):
            self.assertIn(name, names)
        self.assertEqual(next(i['count'] for i in self.scan(text) if i['value'] == '赵星辰'), 2)

    def test_body_business_phrases_are_not_names(self):
        text = '支付金额应当明确。通知费用已结算。双方同意签订合同。公司要求交付。'
        self.assertEqual([], [i['value'] for i in self.scan(text) if i['type'] == 'PERSON_NAME'])

    def test_amount_includes_first_digit(self):
        for value in ('60000元', '6 000元', '6,000.00元', '6000.000元', '6.5万元', '600000元', '660000元'):
            with self.subTest(value=value):
                items = self.scan('应支付金额：' + value + '。')
                self.assertIn(value, [i['value'] for i in items if i['type'] == 'AMOUNT'])

    def test_body_and_amount_export_has_no_prefix_leak(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / 'fictional.docx'
            original = '赵星辰向朱宁支付6 000元。欧阳晓明收到6,000.00元后通知钱晓。'
            doc = Document(); doc.add_paragraph(original); doc.save(src)
            out = Path(tmp) / 'redacted.docx'
            result = R.redact_file(str(src), types=['PERSON_NAME', 'AMOUNT'], out_path=str(out), export_mode='external')
            text = Document(out).paragraphs[0].text
            for value in ('赵星辰', '朱宁', '欧阳晓明', '钱晓', '6'):
                self.assertNotIn(value, text)
            restored = Path(tmp) / 'restored.docx'
            R.restore_file(str(out), mapping_path=result['mapping'], out_path=str(restored))
            self.assertEqual(Document(restored).paragraphs[0].text, original)

    def test_party_names_and_residence(self):
        for label, name in [('原告', '赵星辰'), ('被告', '朱宁'), ('上诉人', '钱晓明'),
                            ('被上诉人', '孙明'), ('被申请人', '李星河'), ('第三人', '吴晓')]:
            with self.subTest(label=label):
                text = f'{label}：{name}，男，汉族，住江苏省示例市石桥镇星光村十二组45号。'
                items = self.scan(text)
                self.assertIn(name, [i['value'] for i in items if i['type'] == 'PERSON_NAME'])
                self.assertIn('江苏省示例市石桥镇星光村十二组45号',
                              [i['value'] for i in items if i['type'] == 'ADDRESS'])

    def test_export_removes_names_and_residence_from_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / 'fictional.docx'
            doc = Document()
            doc.add_paragraph('原告：赵星辰，男，住江苏省示例市石桥镇星光村十二组45号。')
            doc.save(src)
            out = Path(tmp) / 'redacted.docx'
            R.redact_file(str(src), types=['PERSON_NAME', 'ADDRESS'], out_path=str(out))
            with zipfile.ZipFile(out) as z:
                for name in z.namelist():
                    data = z.read(name)
                    self.assertNotIn('赵星辰'.encode(), data)
                    self.assertNotIn('江苏省示例市石桥镇星光村十二组45号'.encode(), data)

    def test_residence_does_not_match_hospital_narrative(self):
        items = self.scan('经治疗后，住院期间应当注意休息。')
        self.assertEqual([], [i for i in items if i['type'] == 'ADDRESS'])

    def test_party_narrative_and_company_are_not_people(self):
        items = self.scan('原告主张解除合同。被告表示同意。原告：北京示例科技有限公司。')
        self.assertEqual([], [i['value'] for i in items if i['type'] == 'PERSON_NAME'])

if __name__ == '__main__':
    unittest.main()
