# -*- coding: utf-8 -*-
import tempfile, unittest, zipfile, sys, json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
import contract_sensitive_detector as D
import contract_redactor as R
class Safety(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def test_tabs_breaks_split_runs_roundtrip(self):
  for sep in ['\t','\n','']:
   with self.subTest(sep=repr(sep)):
    d=Document();p=d.add_paragraph('电话：'+sep);p.add_run('13800');p.add_run('138000');p.add_run('；结束');src=self.root/'in.docx';d.save(src)
    self.assertTrue(all('/' not in b.location for b in D.extract_docx(str(src)).blocks))
    r=R.redact_file(str(src),types=['PHONE_MOBILE'],out_path=str(self.root/'out.docx'),restore_mode=True)
    self.assertEqual(Document(r['output']).paragraphs[0].text,'电话：'+sep+'[手机号码#1]；结束')
    back=R.restore_file(r['output'],out_path=str(self.root/'back.docx'));self.assertEqual(Document(back['output']).paragraphs[0].text,p.text)
 def test_hidden_parts(self):
  d=Document();p=d.add_paragraph('说明');d.add_comment(p.runs,text='电话：13800138000',author='Fixture');s=d.sections[0];s.different_first_page_header_footer=True;s.first_page_header.paragraphs[0].text='电话：13800138000';d.settings.odd_and_even_pages_header_footer=True;s.even_page_footer.paragraphs[0].text='电话：13800138000'
  p=d.add_paragraph();deleted=OxmlElement('w:del');run=OxmlElement('w:r');n=OxmlElement('w:delText');n.text='电话：13800138000';run.append(n);deleted.append(run);p._p.append(deleted);src=self.root/'in.docx';d.save(src)
  r=R.redact_file(str(src),types=['PHONE_MOBILE'],out_path=str(self.root/'out.docx'),export_mode='external');self.assertEqual(r['applied'],4)
  with zipfile.ZipFile(r['output']) as z:
   self.assertFalse(any(b'13800138000' in z.read(n) for n in z.namelist()));self.assertNotIn('customXml/mapping.json',z.namelist())
  self.assertTrue(Path(r['mapping']).is_file());self.assertFalse(r['mapping_embedded']);self.assertIsNone(r['log'])
 def test_repeated_rules_and_redaction(self):
  rules=self.root/'rules.md';rules.write_text('| 28 | REVIEW_CODE_28 | 商业信息 | 中 | 工号 | 隐藏 | keep_head=2; keep_tail=3 | 是 | ZQC[0-9]{8} |\n');src=self.root/'in.txt';src.write_text('员工工号：ZQC58374926')
  for i in range(2):
   specs=D.load_rules_md(str(rules));self.assertEqual(len(specs),1);cs=D.DetectorEngine(D.Settings(extra_specs=specs)).scan_document(D.load_document(str(src)));self.assertTrue(any(c.type_id=='REVIEW_CODE_28' for c in cs));self.assertEqual(next(c.masked for c in cs if c.type_id=='REVIEW_CODE_28'),'ZQ******926');r=R.redact_file(str(src),types=['REVIEW_CODE_28'],rules_path=str(rules),out_path=str(self.root/f'out{i}.txt'));self.assertEqual(r['applied'],1)
 def test_external_removes_mapping_and_blocks_embeddings(self):
  d=Document();d.add_paragraph('电话：13800138000');src=self.root/'in.docx';d.save(src)
  r=R.redact_file(str(src),types=['PHONE_MOBILE'],out_path=str(self.root/'internal.docx'),restore_mode=True);R.embed_mapping_in_docx(r['output'],json.loads(Path(r['mapping']).read_text()));out=R.redact_file(r['output'],types=['PHONE_MOBILE'],out_path=str(self.root/'external.docx'),export_mode='external')
  with zipfile.ZipFile(out['output']) as z:self.assertFalse(any('mapping.json' in n for n in z.namelist()))
  with zipfile.ZipFile(src,'a') as z:z.writestr('word/embeddings/test.bin',b'synthetic')
  with self.assertRaisesRegex(RuntimeError,'嵌入|检查'):R.redact_file(str(src),types=['PHONE_MOBILE'],out_path=str(self.root/'blocked.docx'),export_mode='external')
 def test_input_is_not_overwritten_and_bad_selection_does_not_expand(self):
  d=Document();d.add_paragraph('电话：13800138000');src=self.root/'in.docx';d.save(src);original=src.read_bytes()
  with self.assertRaises(ValueError):R.redact_file(str(src),types=['PHONE_MOBILE'],out_path=str(src))
  self.assertEqual(src.read_bytes(),original)
  with self.assertRaisesRegex(RuntimeError,'复核'):
   R.redact_file(str(src),selection={'enabled_types':['PHONE_MOBILE'],'items':[{'type':'PHONE_MOBILE','offset':[999,1010]}]},out_path=str(self.root/'out.docx'))

if __name__=='__main__':unittest.main()
